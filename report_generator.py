import os
from datetime import datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from models import Downtime, Ship
from kpi_calculator import KPICalculator
from config import Config

class ReportGenerator:
    
    @staticmethod
    def generate_quarterly_report(quarter, year, output_dir='reports'):
        """Генерация отчёта по форме Приложения №6"""
        
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        # Получаем данные КПЭ
        kpi_data = KPICalculator.calculate_kpi_group_a(year, quarter)
        
        # Получаем структуру простоев
        structure = KPICalculator.get_downtime_structure(
            start_date=kpi_data['start_date'],
            end_date=kpi_data['end_date']
        )
        
        # Получаем список простоев за период
        downtimes = Downtime.query.filter(
            Downtime.status == 'approved',
            Downtime.engineer_approved == True,
            Downtime.start_time >= kpi_data['start_date'],
            Downtime.end_time <= kpi_data['end_date']
        ).all()
        
        # Создаём Excel файл
        filename = f"КПЭ5_отчёт_кв{quarter}_{year}.xlsx"
        filepath = os.path.join(output_dir, filename)
        
        wb = Workbook()
        
        # Лист 1: КПЭ №5
        ws1 = wb.active
        ws1.title = "КПЭ №5"
        
        # Стили
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font_white = Font(bold=True, size=12, color="FFFFFF")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Заголовок
        ws1.merge_cells('A1:F1')
        ws1['A1'] = f"Отчёт по показателю КПЭ №5 за {quarter} квартал {year} года"
        ws1['A1'].font = Font(bold=True, size=14)
        ws1['A1'].alignment = Alignment(horizontal='center')
        
        # Информация о периоде
        ws1.cell(row=2, column=1, value="Период расчёта:")
        ws1.cell(row=2, column=2, value=f"{kpi_data['start_date'].strftime('%d.%m.%Y')} - {kpi_data['end_date'].strftime('%d.%m.%Y')}")
        ws1.merge_cells('A2:B2')
        
        # Таблица КПЭ
        headers = ['Группа судов', 'Плановое время (часы)', 'Внеплановое время (часы)', 
                   'Коэффициент готовности (%)', 'Оценка выполнения', 'Статус']
        
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=4, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        
        # Определяем статус и оценку выполнения
        if kpi_data['kg'] >= 85:
            status = "Цель достигнута"
        elif kpi_data['kg'] < 80:
            status = "Критическое отклонение"
        else:
            status = "Зона пропорционального снижения"
        
        data = [
            'A (рабочее ядро)',
            kpi_data['total_planned_hours'],
            kpi_data['total_unplanned_hours'],
            kpi_data['kg'],
            kpi_data['score'],
            status
        ]
        
        for col, value in enumerate(data, 1):
            cell = ws1.cell(row=5, column=col, value=value)
            cell.border = border
            if col == 4:
                if kpi_data['kg'] >= 85:
                    cell.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
                elif kpi_data['kg'] < 80:
                    cell.fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
                else:
                    cell.fill = PatternFill(start_color="FFE4B5", end_color="FFE4B5", fill_type="solid")
        
        # Методика расчёта
        ws1.cell(row=7, column=1, value="Методика расчёта:")
        ws1.cell(row=7, column=1).font = Font(bold=True)
        ws1.cell(row=8, column=1, value="Кг = (Плановое время - Внеплановое время) / Плановое время × 100%")
        ws1.cell(row=9, column=1, value=f"Кг = ({kpi_data['total_planned_hours']} - {kpi_data['total_unplanned_hours']}) / {kpi_data['total_planned_hours']} × 100% = {kpi_data['kg']}%")
        ws1.merge_cells('A7:F7')
        ws1.merge_cells('A8:F8')
        ws1.merge_cells('A9:F9')
        
        # Структура простоев по кодам
        ws1.cell(row=11, column=1, value="Структура простоев по кодам причин")
        ws1.cell(row=11, column=1).font = Font(bold=True, size=12)
        
        structure_headers = ['Код', 'Причина', 'Часы', 'Количество случаев', 'Процент от общего времени']
        for col, header in enumerate(structure_headers, 1):
            cell = ws1.cell(row=12, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.border = border
        
        total_hours = structure.get('_total_hours', 0)
        row = 13
        for code, data in structure.items():
            if code not in ['_total_hours', '_start_date', '_end_date']:
                percent = ((data['hours'] / total_hours) * 100) if total_hours > 0 else 0
                ws1.cell(row=row, column=1, value=code).border = border
                ws1.cell(row=row, column=2, value=data['name']).border = border
                ws1.cell(row=row, column=3, value=round(data['hours'], 2)).border = border
                ws1.cell(row=row, column=4, value=data['count']).border = border
                ws1.cell(row=row, column=5, value=f"{round(percent, 2)}%").border = border
                row += 1
        
        if total_hours > 0:
            ws1.cell(row=row, column=1, value="ИТОГО:").border = border
            ws1.cell(row=row, column=3, value=round(total_hours, 2)).border = border
            ws1.cell(row=row, column=5, value="100%").border = border
        
        # Автоширина колонок
        for col in range(1, ws1.max_column + 1):
            max_length = 0
            column_letter = openpyxl.utils.get_column_letter(col)
            for r in range(1, ws1.max_row + 1):
                cell = ws1.cell(row=r, column=col)
                if isinstance(cell, openpyxl.cell.cell.MergedCell):
                    continue
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws1.column_dimensions[column_letter].width = adjusted_width
        
        # Лист 2: Реестр простоев
        ws2 = wb.create_sheet("Реестр простоев")
        
        reestr_headers = ['ID', 'Судно', 'Группа', 'Начало', 'Окончание', 'Длительность (ч)', 
                         'Код причины', 'Причина', 'Описание', 'Статус', 'Файлы']
        
        for col, header in enumerate(reestr_headers, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.border = border
        
        row = 2
        for dt in downtimes:
            ship = Ship.query.get(dt.ship_id)
            duration = round((dt.end_time - dt.start_time).total_seconds() / 3600, 2) if dt.end_time else 0
            files_list = ', '.join([f.filename for f in dt.files]) if dt.files else ''
            reason_name = Config.DOWNTIME_REASONS.get(dt.reason_code, '')
            
            ws2.cell(row=row, column=1, value=dt.id).border = border
            ws2.cell(row=row, column=2, value=ship.name if ship else '').border = border
            ws2.cell(row=row, column=3, value=ship.group if ship else '').border = border
            ws2.cell(row=row, column=4, value=dt.start_time.strftime('%d.%m.%Y %H:%M')).border = border
            ws2.cell(row=row, column=5, value=dt.end_time.strftime('%d.%m.%Y %H:%M') if dt.end_time else '').border = border
            ws2.cell(row=row, column=6, value=duration).border = border
            ws2.cell(row=row, column=7, value=dt.reason_code).border = border
            ws2.cell(row=row, column=8, value=reason_name).border = border
            ws2.cell(row=row, column=9, value=dt.description[:200] if dt.description else '').border = border
            ws2.cell(row=row, column=10, value=dt.status).border = border
            ws2.cell(row=row, column=11, value=files_list).border = border
            row += 1
        
        # Автоширина для второго листа
        for col in range(1, ws2.max_column + 1):
            max_length = 0
            column_letter = openpyxl.utils.get_column_letter(col)
            for r in range(1, ws2.max_row + 1):
                cell = ws2.cell(row=r, column=col)
                if isinstance(cell, openpyxl.cell.cell.MergedCell):
                    continue
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws2.column_dimensions[column_letter].width = adjusted_width
        
        wb.save(filepath)
        return filepath